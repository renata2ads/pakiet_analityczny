"""Notatka dla zarządu na jedną stronę A4: liczby wyłącznie z ledgera i wyniki/wycena.json.

Użycie: python scripts/generuj_docx.py TICKER
Tworzy TICKER/wyniki/TICKER_notatka_zarzad.docx, zamienia go na PDF i sprawdza, że ma 1 stronę.
"""
import json
import re
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from generuj_pptx import liczba, mld, opis_faktu
from generuj_xlsx import Ledger, soffice
from wspolne import katalog_spolki, wczytaj_ledger, wymagaj_weryfikacji

GRANAT = RGBColor(0x1B, 0x2A, 0x41)
SZARY = RGBColor(0x6B, 0x77, 0x85)
CZCIONKA = "Calibri"
ROZMIAR = 10.5
MIN_ROZMIAR = 9


def run(p, tekst, rozmiar=None, pogrubiony=False, kolor=None):
    r = p.add_run(tekst)
    r.font.name = CZCIONKA
    r.font.size = Pt(rozmiar or ROZMIAR)
    r.font.bold = pogrubiony
    if kolor:
        r.font.color.rgb = kolor
    return r


def ref(p, *idy, rozmiar=None):
    run(p, f" [{', '.join(idy)}]", rozmiar=(rozmiar or ROZMIAR) - 2, kolor=SZARY)


def akapit(doc, po=4, przed=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(po)
    p.paragraph_format.space_before = Pt(przed)
    return p


def sekcja(doc, napis, rozmiar):
    p = akapit(doc, po=2, przed=6)
    run(p, napis, rozmiar=rozmiar + 1.5, pogrubiony=True, kolor=GRANAT)


def punkt(doc, fragmenty, idy, rozmiar):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    for t in fragmenty:
        run(p, t, rozmiar=rozmiar)
    ref(p, *idy, rozmiar=rozmiar)


def cieniowanie(komorka, kolor):
    tc = komorka._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), kolor)
    tc.append(shd)


def zbuduj(ticker, plik, rozmiar):
    L = Ledger(wczytaj_ledger(ticker))
    wynik = katalog_spolki(ticker) / "wyniki"
    wycena = json.loads((wynik / "wycena.json").read_text(encoding="utf-8"))
    info = json.loads((katalog_spolki(ticker) / "zrodla" / "10k" / "info.json").read_text(encoding="utf-8"))
    lata = L.lata("revenue")[-3:]
    ost, prz = lata[-1], lata[-2]
    rev, rev_p = L[f"revenue_{ost}"], L[f"revenue_{prz}"]
    ebit = L[f"operating_income_{ost}"]
    kurs = L["share_price"]

    doc = Document()
    sekcja_ = doc.sections[0]
    sekcja_.page_height, sekcja_.page_width = Cm(29.7), Cm(21.0)
    sekcja_.top_margin = sekcja_.bottom_margin = Cm(1.6)
    sekcja_.left_margin = sekcja_.right_margin = Cm(2.0)
    styl = doc.styles["Normal"]
    styl.font.name = CZCIONKA
    styl.font.size = Pt(rozmiar)

    # Nagłówek
    p = akapit(doc, po=0)
    run(p, "Notatka dla zarządu", rozmiar=rozmiar, pogrubiony=True, kolor=SZARY)
    p = akapit(doc, po=0)
    run(p, info["nazwa_spolki"], rozmiar=20, pogrubiony=True, kolor=GRANAT)
    p = akapit(doc, po=6)
    run(p, f"Data: {date.today().isoformat()}  |  dane: raport 10-K za {ost} r.", rozmiar=rozmiar - 1, kolor=SZARY)
    ref(p, rev["id"], rozmiar=rozmiar - 1)
    run(p, f", kurs z dnia {kurs['okres']}", rozmiar=rozmiar - 1, kolor=SZARY)
    ref(p, kurs["id"], rozmiar=rozmiar - 1)

    # Wnioski
    sekcja(doc, "Najważniejsze wnioski", rozmiar)
    punkt(doc, [f"Przychody wzrosły do {mld(rev)} w {ost} r. z {mld(rev_p)} rok wcześniej, "
                f"a zysk operacyjny wyniósł {mld(ebit)}."], [rev["id"], rev_p["id"], ebit["id"]], rozmiar)
    capex = [w for w in L.z_prefiksem("capex_") if w["klasa"] == "fakt"][-1]
    idy_inw = [capex["id"]]
    tekst_inw = f"Spółka bardzo mocno inwestuje: nakłady w {ost} r. to {liczba(abs(capex['wartosc']) / 1000)} mld USD"
    plan = sorted(k for k in L.po_kluczu if re.search(r"capex_plan_\d{4}_low$", k))
    if plan and plan[-1].replace("_low", "_high") in L.po_kluczu:
        lo, hi = L[plan[-1]], L[plan[-1].replace("_low", "_high")]
        tekst_inw += f", a plan na {lo['okres']} r. to {liczba(lo['wartosc'], 0)}-{liczba(hi['wartosc'], 0)} mld USD"
        idy_inw += [lo["id"], hi["id"]]
    punkt(doc, [tekst_inw + "."], idy_inw, rozmiar)
    roz = wycena["roznica_proc"]
    punkt(doc, [f"Model DCF wycenia akcję na {liczba(wycena['wartosc_jednej_akcji_usd'], 2)} USD wobec kursu "
                f"{liczba(wycena['kurs_usd'], 2)} USD, czyli o {liczba(abs(roz))}% {'niżej' if roz < 0 else 'wyżej'}; "
                "wynik zależy od założeń, a nie od faktów."], ["wycena", kurs["id"]], rozmiar)

    # Tabela kluczowych liczb
    sekcja(doc, "Kluczowe liczby", rozmiar)
    wiersze = [("Pozycja", "Wartość", "id")]
    wiersze.append((f"Przychody {ost}", mld(rev), rev["id"]))
    if f"operating_margin_{ost}" in L.po_kluczu:
        m = L[f"operating_margin_{ost}"]
        wiersze.append((f"Marża operacyjna {ost}", f"{liczba(m['wartosc'], 0)}%", m["id"]))
    gotowka, papiery = L[f"cash_and_equivalents_{ost}"], L[f"marketable_securities_{ost}"]
    wiersze.append(("Gotówka + papiery wartościowe (koniec roku)",
                    f"{liczba(gotowka['wartosc'] / 1000)} + {liczba(papiery['wartosc'] / 1000)} mld USD",
                    f"{gotowka['id']}, {papiery['id']}"))
    wiersze.append(("Wartość jednej akcji z DCF / kurs",
                    f"{liczba(wycena['wartosc_jednej_akcji_usd'], 2)} / {liczba(wycena['kurs_usd'], 2)} USD",
                    f"wycena, {kurs['id']}"))
    tab = doc.add_table(rows=len(wiersze), cols=3)
    tab.alignment = WD_TABLE_ALIGNMENT.LEFT
    tab.style = "Table Grid"
    tab.autofit = False
    for kolumna, szer in zip(tab.columns, (Cm(8.5), Cm(5.5), Cm(3.0))):
        kolumna.width = szer
    for i, wiersz in enumerate(wiersze):
        for j, (tekst, szer) in enumerate(zip(wiersz, (Cm(8.5), Cm(5.5), Cm(3.0)))):
            c = tab.cell(i, j)
            c.width = szer
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            run(c.paragraphs[0], tekst, rozmiar=rozmiar - (1.5 if j == 2 else 0), pogrubiony=i == 0,
                kolor=SZARY if j == 2 and i else None)
            if j == 1 and i:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if i == 0:
                cieniowanie(c, "DCE6F1")

    # Projekty
    sekcja(doc, "Projekty", rozmiar)
    for w in [w for w in L.z_prefiksem("project_") if w["klasa"] == "fakt"
              and not re.search(r"capex_plan_\d{4}_high$", w["klucz"])][:3]:
        punkt(doc, [opis_faktu(w).rstrip(".") + "."], [w["id"]], rozmiar)

    # Ryzyka
    sekcja(doc, "Ryzyka", rozmiar)
    ryzyka = [w for w in L.z_prefiksem("risk_") if w["klasa"] == "fakt"][:3]
    for w in ryzyka:
        punkt(doc, [w["opis"].replace("Ryzyko: ", "").rstrip(".") + "."], [w["id"]], rozmiar)

    # Źródła
    strony = sorted({w["zrodlo"]["strona"] for w in L.wpisy
                     if w["klasa"] == "fakt" and isinstance(w["zrodlo"]["strona"], int)})
    idy_f = [w["id"] for w in L.wpisy if w["klasa"] == "fakt"]
    idy_r = [w["id"] for w in L.wpisy if w["klasa"] == "dana_rynkowa"]
    p = akapit(doc, po=2, przed=8)
    run(p, f"Źródła: raport roczny 10-K {info['nazwa_spolki']} za {ost} r. (SEC EDGAR), strony "
           f"{', '.join(str(s) for s in strony)}; wpisy ledgera {idy_f[0]}-{idy_f[-1]}. "
           f"Dane rynkowe: {', '.join(w['zrodlo']['dokument'] for w in L.wpisy if w['klasa'] == 'dana_rynkowa')} "
           f"({', '.join(idy_r)}). [wycena] = wynik modelu {ticker}_DCF.xlsx (plik wyniki/wycena.json), "
           "oparty na założeniach z ledgera (wpisy Z).", rozmiar=MIN_ROZMIAR, kolor=SZARY)
    p = akapit(doc, po=0)
    run(p, "Materiał edukacyjny, nie stanowi rekomendacji inwestycyjnej.", rozmiar=MIN_ROZMIAR,
        pogrubiony=True, kolor=SZARY)

    doc.save(plik)


def do_pdf(plik):
    import pymupdf
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        profil = Path(tmp) / "profil"
        subprocess.run([soffice(), f"-env:UserInstallation={profil.as_uri()}", "--headless",
                        "--convert-to", "pdf", "--outdir", str(plik.parent), str(plik)],
                       check=True, capture_output=True, timeout=180)
    pdf_plik = plik.with_suffix(".pdf")
    with pymupdf.open(pdf_plik) as pdf:
        podglad = plik.parent / "podglad"
        podglad.mkdir(exist_ok=True)
        for i, strona in enumerate(pdf, start=1):
            strona.get_pixmap(dpi=90).save(podglad / f"docx-{i:02d}.png")
        return len(pdf)


def main():
    if len(sys.argv) != 2:
        sys.exit("Użycie: python scripts/generuj_docx.py TICKER")
    ticker = sys.argv[1].upper()
    wymagaj_weryfikacji(ticker)
    wynik = katalog_spolki(ticker) / "wyniki"
    if not (wynik / "wycena.json").exists():
        sys.exit(f"Brak wyniki/wycena.json. Uruchom najpierw: python scripts/generuj_xlsx.py {ticker}")
    plik = wynik / f"{ticker}_notatka_zarzad.docx"
    for stary in (wynik / "podglad").glob("docx-*.png"):
        stary.unlink()
    rozmiar = ROZMIAR
    while True:
        zbuduj(ticker, plik, rozmiar)
        stron = do_pdf(plik)
        if stron == 1 or rozmiar <= MIN_ROZMIAR:
            break
        rozmiar -= 0.5
    if stron != 1:
        sys.exit(f"Notatka ma {stron} strony nawet przy czcionce {MIN_ROZMIAR} pt. Skróć tekst.")
    print(f"Zapisano {plik} i {plik.with_suffix('.pdf').name}: 1 strona, czcionka {rozmiar} pt")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
