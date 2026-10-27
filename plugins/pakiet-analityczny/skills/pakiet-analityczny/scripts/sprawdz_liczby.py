"""Kontrola liczb: każda liczba w prezentacji i notatce musi mieć pokrycie w ledgerze albo w wycena.json.

Użycie: python scripts/sprawdz_liczby.py TICKER
Pomija lata, daty, numery slajdów, id wpisów, adresy www i nazwy licencji.
Uwzględnia zaokrąglenia oraz zapis w mld (wartość w mln / 1000) i wartość bezwzględną.
"""
import json
import re
import sys

from docx import Document
from pptx import Presentation

from wspolne import katalog_spolki, wczytaj_ledger

WZOR_LICZBY = re.compile(r"\d{1,3}(?:[  ]\d{3})+(?:,\d+)?|\d+(?:,\d+)?")
DO_USUNIECIA = [
    r"https?://\S+",                         # adresy www
    r"\b[FRZG]-\d{2,3}\b",                   # id wpisów ledgera
    r"\b\d{4}-\d{2}-\d{2}\b",                # daty
    r"\b10-K\b", r"\bDGS10\b", r"\bItem 1A\b",
    r"\b\d+-letni\w*",                       # np. 10-letnich (opis obligacji)
    r"\bCC BY(?:-SA)? \d\.\d\b", r"\bCC0\b",  # nazwy licencji
    r"\b[A-Z]{2,}_DCF\b",
]


def wartosci_odniesienia(ticker):
    wartosci = []
    for w in wczytaj_ledger(ticker):
        if isinstance(w.get("wartosc"), (int, float)) and not isinstance(w.get("wartosc"), bool):
            wartosci.append(float(w["wartosc"]))
        strona = (w.get("zrodlo") or {}).get("strona")
        if isinstance(strona, int):
            wartosci.append(float(strona))
    wycena = json.loads((katalog_spolki(ticker) / "wyniki" / "wycena.json").read_text(encoding="utf-8"))
    wartosci += [float(v) for v in wycena.values() if isinstance(v, (int, float))]
    return wartosci


def pokryta(tekst_liczby, wartosci):
    miejsca = len(tekst_liczby.split(",")[1]) if "," in tekst_liczby else 0
    x = float(re.sub(r"[  ]", "", tekst_liczby).replace(",", "."))
    for v in wartosci:
        for kandydat in (v, v / 1000, v / 1_000_000):
            if abs(round(abs(kandydat), miejsca) - x) < 1e-9:
                return True
    return False


def czy_rok(tekst_liczby):
    return re.fullmatch(r"(19|20)\d{2}", tekst_liczby) is not None


def teksty_pptx(plik):
    for nr, slajd in enumerate(Presentation(plik).slides, start=1):
        for ksztalt in slajd.shapes:
            if ksztalt.has_text_frame:
                for p in ksztalt.text_frame.paragraphs:
                    yield f"slajd {nr}", "".join(r.text for r in p.runs)


def teksty_docx(plik):
    doc = Document(plik)
    for nr, p in enumerate(doc.paragraphs, start=1):
        yield f"akapit {nr}", p.text
    for t, tabela in enumerate(doc.tables, start=1):
        for i, wiersz in enumerate(tabela.rows, start=1):
            for j, komorka in enumerate(wiersz.cells, start=1):
                yield f"tabela {t}, wiersz {i}, kolumna {j}", komorka.text


def main():
    if len(sys.argv) != 2:
        sys.exit("Użycie: python scripts/sprawdz_liczby.py TICKER")
    ticker = sys.argv[1].upper()
    wynik = katalog_spolki(ticker) / "wyniki"
    wartosci = wartosci_odniesienia(ticker)
    pliki = [(wynik / f"{ticker}_prezentacja.pptx", teksty_pptx),
             (wynik / f"{ticker}_notatka_zarzad.docx", teksty_docx)]
    bledy, sprawdzone = [], 0
    for plik, czytnik in pliki:
        if not plik.exists():
            bledy.append(f"{plik.name}: brak pliku")
            continue
        for miejsce, tekst in czytnik(plik):
            oczyszczony = tekst
            for wzor in DO_USUNIECIA:
                oczyszczony = re.sub(wzor, " ", oczyszczony)
            for m in WZOR_LICZBY.finditer(oczyszczony):
                liczba = m.group(0).strip()
                if czy_rok(liczba):
                    continue
                sprawdzone += 1
                if not pokryta(liczba, wartosci):
                    bledy.append(f"{plik.name}, {miejsce}: liczba {liczba!r} bez pokrycia w ledgerze "
                                 f"ani w wycena.json | {tekst.strip()[:90]}")
    print(f"Sprawdzono liczb: {sprawdzone}")
    if bledy:
        print(f"BŁĘDY ({len(bledy)}):")
        for b in bledy:
            print("  - " + b)
        sys.exit(1)
    print("OK: każda liczba ma pokrycie w ledgerze albo w wycena.json.")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
