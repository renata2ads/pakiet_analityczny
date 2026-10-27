"""Prezentacja o spółce: liczby i fakty wyłącznie z ledgera i wyniki/wycena.json.

Użycie: python scripts/generuj_pptx.py TICKER
Tworzy TICKER/wyniki/TICKER_prezentacja.pptx i podgląd slajdów w TICKER/wyniki/podglad/.
"""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from generuj_xlsx import Ledger, soffice
from wspolne import KATALOG_GLOWNY, katalog_spolki, wczytaj_ledger, wymagaj_weryfikacji

GRANAT = RGBColor(0x1B, 0x2A, 0x41)
MORSKI = RGBColor(0x0F, 0x7C, 0x8C)
JASNY = RGBColor(0xEE, 0xF3, 0xF7)
BIALY = RGBColor(0xFF, 0xFF, 0xFF)
TEKST = RGBColor(0x22, 0x2B, 0x36)
SZARY = RGBColor(0x6B, 0x77, 0x85)
OSTRZEZENIE = RGBColor(0xFD, 0xEC, 0xD2)
POMARANCZ = RGBColor(0xB4, 0x5F, 0x06)
NAGLOWEK = "Cambria"
TRESC = "Calibri"

SZER, WYS = Inches(13.333), Inches(7.5)
MARGINES = Inches(0.6)


def liczba(x, miejsca=1):
    """Polski zapis liczby: spacja jako separator tysięcy, przecinek dziesiętny."""
    tekst = f"{x:,.{miejsca}f}".replace(",", " ").replace(".", ",")
    return tekst


def mld(w):
    """Kwota z ledgera w mln USD jako tekst w mld USD."""
    return f"{liczba(w['wartosc'] / 1000)} mld USD"


def tekst(slajd, x, y, w, h, akapity, rozmiar=16, kolor=TEKST, czcionka=TRESC, pogrubiony=False,
          wyrownanie=PP_ALIGN.LEFT, kotwica=MSO_ANCHOR.TOP, odstep=6):
    """Pole tekstowe. akapity: lista; element to str albo lista fragmentów (tekst, {opcje})."""
    pole = slajd.shapes.add_textbox(x, y, w, h)
    ramka = pole.text_frame
    ramka.word_wrap = True
    ramka.vertical_anchor = kotwica
    ramka.margin_left = ramka.margin_right = ramka.margin_top = ramka.margin_bottom = 0
    for i, akapit in enumerate(akapity if isinstance(akapity, list) else [akapity]):
        p = ramka.paragraphs[0] if i == 0 else ramka.add_paragraph()
        p.alignment = wyrownanie
        p.space_after = Pt(odstep)
        for fragment, opcje in (akapit if isinstance(akapit, list) else [(akapit, {})]):
            r = p.add_run()
            r.text = fragment
            f = r.font
            f.name = opcje.get("czcionka", czcionka)
            f.size = Pt(opcje.get("rozmiar", rozmiar))
            f.bold = opcje.get("pogrubiony", pogrubiony)
            f.color.rgb = opcje.get("kolor", kolor)
    return pole


def ref(*idy):
    return (f"  [{', '.join(idy)}]", {"rozmiar": 10, "kolor": SZARY, "pogrubiony": False})


def tlo(slajd, kolor):
    slajd.background.fill.solid()
    slajd.background.fill.fore_color.rgb = kolor


def prostokat(slajd, x, y, w, h, kolor, zaokraglony=True):
    ksztalt = slajd.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if zaokraglony else MSO_SHAPE.RECTANGLE,
                                     x, y, w, h)
    ksztalt.fill.solid()
    ksztalt.fill.fore_color.rgb = kolor
    ksztalt.line.fill.background()
    ksztalt.shadow.inherit = False
    if zaokraglony:
        ksztalt.adjustments[0] = 0.08
    return ksztalt


def tytul(slajd, napis):
    tekst(slajd, MARGINES, Inches(0.45), SZER - 2 * MARGINES, Inches(0.8), napis,
          rozmiar=34, kolor=GRANAT, czcionka=NAGLOWEK, pogrubiony=True)


def stopka(slajd, napis):
    tekst(slajd, MARGINES, WYS - Inches(0.5), SZER - 2 * MARGINES, Inches(0.3), napis, rozmiar=10, kolor=SZARY)


def zbuduj(ticker, plik):
    L = Ledger(wczytaj_ledger(ticker))
    wynik = katalog_spolki(ticker) / "wyniki"
    wycena = json.loads((wynik / "wycena.json").read_text(encoding="utf-8"))
    info = json.loads((katalog_spolki(ticker) / "zrodla" / "10k" / "info.json").read_text(encoding="utf-8"))
    lata = L.lata("revenue")[-3:]
    ost = lata[-1]
    rev = {r: L[f"revenue_{r}"] for r in lata}
    ebit = {r: L[f"operating_income_{r}"] for r in lata}

    prs = Presentation()
    prs.slide_width, prs.slide_height = SZER, WYS
    pusty = prs.slide_layouts[6]

    # 1. Tytuł
    s = prs.slides.add_slide(pusty)
    tlo(s, GRANAT)
    tekst(s, MARGINES, Inches(2.2), Inches(11), Inches(1.2), info["nazwa_spolki"],
          rozmiar=48, kolor=BIALY, czcionka=NAGLOWEK, pogrubiony=True)
    tekst(s, MARGINES, Inches(3.4), Inches(11), Inches(0.7), "Pakiet analityczny",
          rozmiar=28, kolor=RGBColor(0x9F, 0xD8, 0xE0))
    kurs = L["share_price"]
    tekst(s, MARGINES, Inches(4.5), Inches(11), Inches(1.0), [
        [(f"Dane finansowe: raport roczny 10-K za {ost} r. (SEC)", {}), ref(rev[ost]["id"])],
        [(f"Dane rynkowe: kurs akcji z dnia {kurs['okres']}", {}), ref(kurs["id"])],
    ], rozmiar=16, kolor=BIALY)
    s.notes_slide.notes_text_frame.text = "Każda liczba w prezentacji ma id wpisu ledgera (F - fakt, R - dana rynkowa, Z - założenie, G - grafika)."

    # 2. Spółka w liczbach
    s = prs.slides.add_slide(pusty)
    tlo(s, BIALY)
    tytul(s, "Spółka w liczbach")
    dane = CategoryChartData()
    dane.categories = [str(r) for r in lata]
    dane.add_series("Przychody", [rev[r]["wartosc"] / 1000 for r in lata])
    dane.add_series("Zysk operacyjny", [ebit[r]["wartosc"] / 1000 for r in lata])
    wykres = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, MARGINES, Inches(1.5),
                                Inches(7.6), Inches(4.9), dane).chart
    wykres.has_title = True
    wykres.chart_title.text_frame.text = "Przychody i zysk operacyjny (mld USD)"
    wykres.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    wykres.chart_title.text_frame.paragraphs[0].runs[0].font.name = TRESC
    wykres.has_legend = True
    wykres.legend.position = XL_LEGEND_POSITION.BOTTOM
    wykres.legend.include_in_layout = False
    wykres.legend.font.size = Pt(12)
    wykres.value_axis.has_major_gridlines = True
    wykres.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xE3, 0xEA)
    wykres.value_axis.tick_labels.font.size = Pt(11)
    wykres.value_axis.format.line.fill.background()
    wykres.category_axis.tick_labels.font.size = Pt(12)
    for seria, kolor in zip(wykres.plots[0].series, (GRANAT, MORSKI)):
        seria.format.fill.solid()
        seria.format.fill.fore_color.rgb = kolor
    wykres.plots[0].has_data_labels = True
    etyk = wykres.plots[0].data_labels
    etyk.number_format = '0.0'
    etyk.number_format_is_linked = False
    etyk.position = XL_LABEL_POSITION.OUTSIDE_END
    etyk.font.size = Pt(11)
    idy_wykresu = [rev[r]["id"] for r in lata] + [ebit[r]["id"] for r in lata]
    tekst(s, MARGINES, Inches(6.45), Inches(7.6), Inches(0.3),
          [[(f"Źródło: raport 10-K, strona {rev[ost]['zrodlo']['strona']}", {"rozmiar": 10, "kolor": SZARY}),
            ref(*sorted(idy_wykresu))]], rozmiar=10)
    karty = [(mld(rev[ost]), f"przychody w {ost} r.", rev[ost]["id"]),
             (mld(L[f"net_income_{ost}"]), f"zysk netto w {ost} r.", L[f"net_income_{ost}"]["id"])]
    # Trzecia karta: pierwszy liczbowy fakt o projekcie (np. liczba użytkowników), jeśli jest.
    liczbowe = [w for w in L.z_prefiksem("project_") if w["klasa"] == "fakt"
                and isinstance(w["wartosc"], (int, float)) and "capex_plan" not in w["klucz"]]
    if liczbowe:
        tekst_, _, wartosc_ = opis_faktu(liczbowe[0]).rpartition(": ")
        karty.append((wartosc_, tekst_[:60], liczbowe[0]["id"]))
    for i, (duza, opis, id_) in enumerate(karty):
        y = Inches(1.5) + i * Inches(1.7)
        prostokat(s, Inches(8.7), y, Inches(4.0), Inches(1.45), JASNY)
        tekst(s, Inches(8.95), y + Inches(0.15), Inches(3.6), Inches(0.7), duza,
              rozmiar=30, kolor=GRANAT, czcionka=NAGLOWEK, pogrubiony=True)
        tekst(s, Inches(8.95), y + Inches(0.85), Inches(3.6), Inches(0.5), [[(opis, {}), ref(id_)]], rozmiar=14)

    # 3+. Projekty ze zdjęciami
    grafiki = [w for w in L.wpisy if w["klasa"] == "grafika"]
    po_id = {w["id"]: w for w in L.wpisy}
    punkty_projektow = projekty(L, po_id)
    for g in grafiki:
        s = prs.slides.add_slide(pusty)
        tlo(s, BIALY)
        naglowek_, punkty = punkty_projektow.get(g["projekt"], (g["opis"], [(po_id[g["projekt"]]["opis"], [g["projekt"]])]))
        tytul(s, naglowek_)
        sciezka = KATALOG_GLOWNY / g["plik"]
        from PIL import Image
        with Image.open(sciezka) as obraz:
            proporcja = obraz.height / obraz.width
        max_w, max_h = Inches(6.4), Inches(4.6)
        w = max_w
        h = int(w * proporcja)
        if h > max_h:
            h = max_h
            w = int(h / proporcja)
        s.shapes.add_picture(str(sciezka), MARGINES, Inches(1.5), w, h)
        tekst(s, MARGINES, Inches(1.5) + h + Inches(0.1), Inches(6.4), Inches(0.5),
              [[(f"Fot. {g['autor']}, licencja {g['licencja']}, Wikimedia Commons", {}), ref(g["id"])]],
              rozmiar=10, kolor=SZARY)
        prostokat(s, Inches(7.4), Inches(1.5), Inches(5.3), Inches(4.6), JASNY)
        tekst(s, Inches(7.7), Inches(1.8), Inches(4.7), Inches(4.1),
              [[("•  " + t, {}), ref(*idy)] for t, idy in punkty], rozmiar=17, odstep=16)
        stopka(s, f"Źródło faktów: raport 10-K za {ost} r. (SEC)")

    # Wycena DCF
    s = prs.slides.add_slide(pusty)
    tlo(s, BIALY)
    tytul(s, "Wycena DCF: wartość jednej akcji wobec kursu")
    for i, (etykieta, wartosc, kolor) in enumerate([
            ("Wartość jednej akcji z modelu DCF", f"{liczba(wycena['wartosc_jednej_akcji_usd'], 2)} USD", GRANAT),
            ("Kurs akcji na giełdzie", f"{liczba(wycena['kurs_usd'], 2)} USD", MORSKI)]):
        x = MARGINES + i * Inches(3.55)
        prostokat(s, x, Inches(1.5), Inches(3.3), Inches(1.9), JASNY)
        tekst(s, x + Inches(0.25), Inches(1.7), Inches(2.9), Inches(0.8), wartosc,
              rozmiar=32, kolor=kolor, czcionka=NAGLOWEK, pogrubiony=True)
        tekst(s, x + Inches(0.25), Inches(2.55), Inches(2.9), Inches(0.7),
              [[(etykieta, {}), ref("wycena" if i == 0 else kurs["id"])]], rozmiar=14)
    roznica = wycena["roznica_proc"]
    tekst(s, MARGINES, Inches(3.65), Inches(6.9), Inches(0.9),
          [[(f"Model wycenia akcję o {liczba(abs(roznica))}% {'niżej' if roznica < 0 else 'wyżej'} niż giełda.", {}),
            ref("wycena")]], rozmiar=18, pogrubiony=True, kolor=GRANAT)
    zal = [(f"WACC (koszt kapitału): {liczba(wycena['wacc_proc'], 2)}%", ["wycena"]),
           (f"Wzrost przychodów {L.lata('revenue_growth')[0]}-{L.lata('revenue_growth')[-1]}: "
            + " / ".join(liczba(w['wartosc'], 0) for w in L.z_prefiksem("revenue_growth_")) + "%",
            [w["id"] for w in L.z_prefiksem("revenue_growth_")][::len(L.z_prefiksem("revenue_growth_")) - 1]),
           (f"Marża operacyjna: {liczba(L['operating_margin']['wartosc'], 0)}%", [L["operating_margin"]["id"]]),
           (f"Wzrost po {L.lata('revenue_growth')[-1]} r.: {liczba(L['terminal_growth']['wartosc'], 0)}%",
            [L["terminal_growth"]["id"]]),
           (f"Beta {liczba(L['beta']['wartosc'], 1)}, premia za ryzyko {liczba(L['equity_risk_premium']['wartosc'], 1)}%",
            [L["beta"]["id"], L["equity_risk_premium"]["id"]])]
    tekst(s, MARGINES, Inches(4.5), Inches(6.9), Inches(2.3),
          [[("Kluczowe założenia", {"pogrubiony": True, "kolor": GRANAT})]] +
          [[("•  " + t, {}), ref(*idy)] for t, idy in zal], rozmiar=15, odstep=4)
    prostokat(s, Inches(7.9), Inches(1.5), Inches(4.8), Inches(4.9), OSTRZEZENIE)
    tekst(s, Inches(8.2), Inches(1.75), Inches(4.2), Inches(4.4), [
        [("To są założenia, nie fakty", {"rozmiar": 24, "pogrubiony": True, "kolor": POMARANCZ,
                                         "czcionka": NAGLOWEK})],
        "Wynik DCF zależy od przyjętych założeń o przyszłości: tempie wzrostu, marży, nakładach inwestycyjnych i koszcie kapitału.",
        "Niewielka zmiana WACC lub wzrostu rezydualnego wyraźnie zmienia wartość akcji (arkusz Wrażliwość w pliku Excel).",
        "Materiał edukacyjny, nie stanowi rekomendacji inwestycyjnej.",
    ], rozmiar=16, odstep=12)
    stopka(s, f"Model: {ticker}_DCF.xlsx; wynik z pliku wyniki/wycena.json [wycena]")

    # Ryzyka i źródła
    s = prs.slides.add_slide(pusty)
    tlo(s, BIALY)
    tytul(s, "Ryzyka i źródła")
    ryzyka = [w for w in L.z_prefiksem("risk_") if w["klasa"] == "fakt"]
    tekst(s, MARGINES, Inches(1.45), Inches(6.2), Inches(0.5), "Najważniejsze ryzyka (Item 1A, Risk Factors)",
          rozmiar=18, pogrubiony=True, kolor=GRANAT)
    for i, w in enumerate(ryzyka):
        y = Inches(2.05) + i * Inches(1.45)
        prostokat(s, MARGINES, y, Inches(6.2), Inches(1.25), JASNY)
        tekst(s, MARGINES + Inches(0.25), y + Inches(0.15), Inches(5.7), Inches(1.0),
              [[((lambda t: t[:1].upper() + t[1:])(w["opis"].replace("Ryzyko: ", "")), {"pogrubiony": True}), ref(w["id"])],
               [(f"raport 10-K, strona {w['zrodlo']['strona']}", {"rozmiar": 12, "kolor": SZARY})]],
              rozmiar=16, odstep=2, kotwica=MSO_ANCHOR.MIDDLE)
    strony = sorted({w["zrodlo"]["strona"] for w in L.wpisy
                     if w["klasa"] == "fakt" and isinstance(w["zrodlo"]["strona"], int)})
    okladka = any(w["klasa"] == "fakt" and not isinstance(w["zrodlo"]["strona"], int) for w in L.wpisy)
    rynek = [w for w in L.wpisy if w["klasa"] == "dana_rynkowa"]
    zrodla = [
        [(f"Raport roczny 10-K {info['nazwa_spolki']} za {ost} r. (SEC EDGAR)", {"pogrubiony": True})],
        [(f"Strony: {', '.join(str(x) for x in strony)}" + (" oraz okładka" if okladka else ""), {"rozmiar": 13})],
        [(info["url"], {"rozmiar": 10, "kolor": SZARY})],
    ]
    for w in rynek:
        zrodla.append([(w["zrodlo"]["dokument"], {"pogrubiony": True}), ref(w["id"])])
        zrodla.append([(f"pobrano {w['data_pobrania']}", {"rozmiar": 12, "kolor": SZARY})])
    zrodla.append([("Wikimedia Commons: zdjęcia na wolnych licencjach", {"pogrubiony": True}),
                   ref(*[g["id"] for g in grafiki])])
    zrodla.append([("Założenia modelu: analityk, z uzasadnieniem w ledgerze", {"pogrubiony": True}), ref("Z")])
    tekst(s, Inches(7.3), Inches(1.45), Inches(5.4), Inches(0.5), "Źródła", rozmiar=18, pogrubiony=True, kolor=GRANAT)
    tekst(s, Inches(7.3), Inches(2.05), Inches(5.4), Inches(4.6), zrodla, rozmiar=14, odstep=3)
    stopka(s, "Pełna lista wpisów z cytatami: arkusz Ledger w pliku Excel oraz plik ledger.json")

    prs.save(plik)
    return len(prs.slides)


def opis_faktu(w):
    """Zdanie o fakcie: opis z ledgera i wartość w czytelnej jednostce."""
    v, jedn = w.get("wartosc"), w.get("jednostka", "")
    if not isinstance(v, (int, float)) or isinstance(v, bool):
        return w["opis"]
    if jedn == "mln USD" and abs(v) >= 1000:
        tekst_v = f"{liczba(abs(v) / 1000)} mld USD"
    elif jedn == "%":
        tekst_v = f"{liczba(v, 0 if float(v).is_integer() else 1)}%"
    else:
        miejsca = 0 if float(v).is_integer() else len(repr(float(v)).split(".")[1])
        tekst_v = f"{liczba(abs(v) if jedn.endswith('USD') else v, miejsca)} {jedn}".strip()
    return f"{w['opis']}: {tekst_v}"


def projekty(L, po_id):
    """Treść slajdów projektów, z wpisów klasy grafika.

    Wpis grafiki może mieć pola "naglowek" (nazwa projektu po polsku) i "fakty" (lista id faktów
    do punktów slajdu). Bez nich: nagłówek = opis zdjęcia, punkt = fakt z pola "projekt".
    Klucz słownika = id faktu z pola "projekt".
    """
    tresc = {}
    for g in [w for w in L.wpisy if w["klasa"] == "grafika"]:
        fakty = [i for i in (g.get("fakty") or [g["projekt"]]) if i in po_id]
        tresc[g["projekt"]] = (g.get("naglowek") or g["opis"], [(opis_faktu(po_id[i]), [i]) for i in fakty])
    return tresc


def podglad(plik, katalog):
    import pymupdf
    katalog.mkdir(parents=True, exist_ok=True)
    for stary in katalog.glob("pptx-*.png"):
        stary.unlink()
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        profil = Path(tmp) / "profil"
        subprocess.run([soffice(), f"-env:UserInstallation={profil.as_uri()}", "--headless",
                        "--convert-to", "pdf", "--outdir", tmp, str(plik)],
                       check=True, capture_output=True, timeout=180)
        with pymupdf.open(Path(tmp) / (plik.stem + ".pdf")) as pdf:
            for i, strona in enumerate(pdf, start=1):
                strona.get_pixmap(dpi=80).save(katalog / f"pptx-{i:02d}.png")
            return len(pdf)


def main():
    if len(sys.argv) != 2:
        sys.exit("Użycie: python scripts/generuj_pptx.py TICKER")
    ticker = sys.argv[1].upper()
    wymagaj_weryfikacji(ticker)
    wynik = katalog_spolki(ticker) / "wyniki"
    if not (wynik / "wycena.json").exists():
        sys.exit(f"Brak wyniki/wycena.json. Uruchom najpierw: python scripts/generuj_xlsx.py {ticker}")
    plik = wynik / f"{ticker}_prezentacja.pptx"
    n = zbuduj(ticker, plik)
    podglad(plik, wynik / "podglad")
    print(f"Zapisano {plik} ({n} slajdów), podgląd w {wynik / 'podglad'}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
